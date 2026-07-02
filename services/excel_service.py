"""Excel取込・出力ビジネスロジック"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from config import Config
from services.employee_service import create_employee
from database import get_db


def _read_table(file_path: str) -> pd.DataFrame:
    if file_path.lower().endswith(".csv"):
        return pd.read_csv(file_path)
    return pd.read_excel(file_path)


def import_employee_excel(file_path: str) -> tuple[int, list[str]]:
    errors: list[str] = []
    count = 0
    df = _read_table(file_path)
    df = df.rename(columns={
        df.columns[0]: "employee_no",
        df.columns[1]: "name",
        df.columns[2]: "department",
        df.columns[3]: "position",
        df.columns[4]: "hire_date",
    })
    for i, row in df.iterrows():
        try:
            create_employee({
                "employee_no": str(row.get("employee_no", "")).strip(),
                "name": str(row.get("name", "")).strip(),
                "department": str(row.get("department", "")).strip(),
                "position": str(row.get("position", "")).strip(),
                "hire_date": str(row.get("hire_date", ""))[:10] if pd.notna(row.get("hire_date")) else "",
            })
            count += 1
        except Exception as exc:
            errors.append(f"{i + 2}行目: {exc}")
    return count, errors


def import_training_excel(file_path: str, training_id: int) -> tuple[int, list[str]]:
    errors: list[str] = []
    count = 0
    df = _read_table(file_path)
    df = df.rename(columns={
        df.columns[0]: "employee_no",
        df.columns[2]: "attendance_rate",
        df.columns[3]: "understanding_level",
        df.columns[4]: "report_score",
    })
    conn = get_db()
    for i, row in df.iterrows():
        try:
            emp = conn.execute("SELECT id FROM employee WHERE employee_no = ?", (str(row["employee_no"]).strip(),)).fetchone()
            if not emp:
                raise ValueError("社員番号が見つかりません")
            conn.execute(
                """
                INSERT INTO training_history (employee_id, training_id, attendance_rate, understanding_level, report_score)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(employee_id, training_id) DO UPDATE SET
                    attendance_rate = excluded.attendance_rate,
                    understanding_level = excluded.understanding_level,
                    report_score = excluded.report_score
                """,
                (emp["id"], training_id, float(row["attendance_rate"]), int(row["understanding_level"]), int(row["report_score"])),
            )
            count += 1
        except Exception as exc:
            errors.append(f"{i + 2}行目: {exc}")
    conn.commit()
    conn.close()
    return count, errors


def export_evaluation_excel(evaluations: list[dict], output_path: str) -> str:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "評価一覧"
    headers = ["社員番号", "氏名", "部署", "出席率", "理解度", "課題", "総合スコア", "AIコメント"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for row in evaluations:
        ws.append([
            row.get("employee_no"), row.get("name"), row.get("department"),
            row.get("attendance_rate"), row.get("understanding_level"), row.get("report_score"),
            row.get("score") or row.get("calculated_score"), row.get("ai_comment"),
        ])
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max_len + 2, 60)
    wb.save(output_path)
    return output_path


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in Config.ALLOWED_EXTENSIONS
